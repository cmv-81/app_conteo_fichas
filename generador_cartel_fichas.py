#!/usr/bin/env python3
"""
Generador de cartel de cambio de dinero a fichas.

Uso:
    pip install openpyxl matplotlib numpy
    python generador_cartel_fichas.py [plantilla_cambio_fichas.xlsx] [cartel_cambio_fichas.png] [--forzar]

ARQUITECTURA:
    EL EXCEL CALCULA → EL EXCEL GUARDA LOS DATOS → PYTHON LEE LOS DATOS → PYTHON DIBUJA EL CARTEL.

Este programa NO calcula combinaciones ni devoluciones. Lee la hoja DATOS_PYTHON
del Excel (que se rellena sola con fórmulas a partir de CONFIGURACIÓN y CAMBIOS)
y dibuja exactamente:
    - el número de fichas de cada tipo que pone el Excel,
    - la devolución calculada por el Excel,
    - el estado de cada casilla (OK / REVISAR / NO ES POSIBLE).

Si alguna casilla está en REVISAR, el cartel no se genera (para no imprimir
cuentas erróneas). Con --forzar se genera igualmente con los datos del Excel.
"""

import logging
import sys
from pathlib import Path

try:
    import openpyxl
    import matplotlib.pyplot as plt
    import numpy as np
    from matplotlib.patches import Ellipse, Rectangle
except ModuleNotFoundError as e:
    sys.exit(
        f"Falta la librería '{e.name}' en este Python ({sys.executable}).\n"
        f'Instálala con:\n    "{sys.executable}" -m pip install openpyxl matplotlib numpy'
    )

# DejaVu no tiene peso 900 y usa negrita (700); se silencia ese aviso repetido.
logging.getLogger("matplotlib.font_manager").setLevel(logging.ERROR)


# ---------- PALETA DEL CARTEL ----------
DARK = "#082E35"
DARK2 = "#123E45"
CREAM = "#F1EEE4"
CELL = "#EEEFEA"
GREY = "#D8DDD8"
BLACK = "#07161A"
WHITE = "#FFFDF6"
RED = "#B92F28"

# Tamaño del cartel (pulgadas). El eje va de 0 a 1 en ambos sentidos, así que
# para que las fichas salgan redondas el radio horizontal se corrige con ASPECT.
FIG_W, FIG_H = 16, 10.4
ASPECT = FIG_H / FIG_W

# Columnas del cartel: qué tipos de ficha (0, 1, 2) entran en cada una.
COMBINACIONES = ((0,), (1,), (2,), (0, 1), (0, 2), (1, 2))

HOJA_DATOS = "DATOS_PYTHON"
ESTADOS = ("OK", "REVISAR", "NO ES POSIBLE")


# ---------- EXCEL ----------

class ErrorExcel(Exception):
    pass


def leer_excel(path):
    """
    Lee la hoja DATOS_PYTHON:
      - A:B  → pares clave / valor (fichas, bebidas, colores y textos).
      - D:K  → una fila por casilla del cartel:
               Importe | Columna | Combinación | Fichas tipo 1 | Fichas tipo 2 |
               Fichas tipo 3 | Devolución | Estado
    No calcula nada: devuelve los valores tal cual los guarda el Excel.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    if HOJA_DATOS not in wb.sheetnames:
        raise ErrorExcel(f"El Excel no tiene la hoja '{HOJA_DATOS}'.")
    ws = wb[HOJA_DATOS]

    # --- Claves ---
    claves = {}
    for r in range(5, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if k is None:
            break
        claves[str(k).strip()] = ws.cell(r, 2).value

    sin_valor = [k for k, v in claves.items() if v is None]
    if "ficha1_valor" not in claves or sin_valor:
        raise ErrorExcel(
            "El Excel no tiene los resultados de las fórmulas guardados.\n"
            "Ábrelo en Excel, pulsa Guardar y vuelve a ejecutar el programa."
        )

    valores = [float(claves[f"ficha{i}_valor"]) for i in (1, 2, 3)]
    bebidas = [str(claves[f"ficha{i}_bebida"]) for i in (1, 2, 3)]
    colores = [str(claves[f"ficha{i}_color"]).strip() for i in (1, 2, 3)]
    textos = {
        "titulo": claves.get("titulo") or "CAMBIO DE DINERO A FICHAS",
        "subtitulo": claves.get("subtitulo") or "TÚ ELIGES  •  DISFRUTA DEL FESTIVAL",
        "tabla": claves.get("titulo_tabla") or "POSIBLES COMBINACIONES DE CAMBIO DE FICHAS",
        "pie": claves.get("pie") or "¡ELIGE TU COMBINACIÓN, CAMBIA TUS EUROS POR FICHAS Y DISFRUTA!",
    }

    # --- Tabla de casillas ---
    importes = []
    tabla = {}  # importe -> [casilla columna 1..6]
    for r in range(5, ws.max_row + 1):
        importe = ws.cell(r, 4).value
        if importe is None:
            break
        columna = ws.cell(r, 5).value
        fichas = [ws.cell(r, c).value for c in (7, 8, 9)]
        devolucion = ws.cell(r, 10).value
        estado = ws.cell(r, 11).value

        if columna is None or estado is None or any(f is None for f in fichas):
            raise ErrorExcel(
                f"Fila {r} de {HOJA_DATOS} sin valores calculados.\n"
                "Ábrelo en Excel, pulsa Guardar y vuelve a ejecutar el programa."
            )
        estado = str(estado).strip().upper()
        if estado not in ESTADOS:
            raise ErrorExcel(f"Estado desconocido '{estado}' en la fila {r} de {HOJA_DATOS}.")

        importe = float(importe)
        if importe not in tabla:
            importes.append(importe)
            tabla[importe] = [None] * 6
        tabla[importe][int(columna) - 1] = {
            "estado": estado,
            "fichas": [int(round(float(f))) for f in fichas],
            "devolucion": devolucion,
            "combinacion": ws.cell(r, 6).value,
        }

    for importe in importes:
        if any(c is None for c in tabla[importe]):
            raise ErrorExcel(f"Faltan columnas para el importe {importe:g}€ en {HOJA_DATOS}.")

    return valores, bebidas, colores, importes, tabla, textos


def casillas_a_revisar(tabla, importes):
    return [
        (importe, tabla[importe][c]["combinacion"])
        for importe in importes
        for c in range(6)
        if tabla[importe][c]["estado"] == "REVISAR"
    ]


def fmt_euros(v):
    """Formatea un número del Excel sin decimales sobrantes (5.0 -> '5')."""
    if isinstance(v, (int, float)):
        return f"{float(v):g}"
    return str(v)


# ---------- DIBUJO ----------

def moneda(ax, x, y, valor, color, radio=0.024, zbase=20):
    """Ficha compacta y legible: sombra, borde, cuerpo y pequeño brillo."""
    rx = radio * ASPECT  # radio horizontal en coordenadas del eje → ficha redonda

    def disco(cx, cy, k, **kw):
        ax.add_patch(Ellipse((cx, cy), 2*rx*k, 2*radio*k, **kw))

    disco(x + rx*0.10, y - radio*0.12, 1.05,
          facecolor="#000000", edgecolor="none", alpha=0.16, zorder=zbase)
    disco(x, y, 1.03,
          facecolor="#151515", edgecolor="#151515", linewidth=0.6, zorder=zbase+1)
    disco(x, y, 0.94,
          facecolor=color, edgecolor="#111111", linewidth=0.65, zorder=zbase+2)
    disco(x - rx*0.27, y + radio*0.27, 0.19,
          facecolor="#FFFFFF", edgecolor="none", alpha=0.20, zorder=zbase+3)
    ax.text(
        x, y, f"{valor:g}€",
        ha="center", va="center",
        fontsize=max(6.5, radio*390),
        fontweight="900", color=BLACK,
        zorder=zbase+4
    )


def ancho_texto(ax, artista):
    """Ancho real (en coordenadas del eje) de un texto ya dibujado."""
    renderer = ax.figure.canvas.get_renderer()
    bb = artista.get_window_extent(renderer=renderer)
    return bb.width / ax.figure.bbox.width


def cantidad_y_ficha(ax, cx, cy, cantidad, valor, color, fontsize, radio, hueco, ancho_max):
    """
    Dibuja «NÚMERO [ficha]» centrado en cx.
    Se mide el número real para que la ficha empiece SIEMPRE después de él
    (nunca lo tapa). Si el grupo no cabe en ancho_max, se reduce la letra.
    """
    while True:
        t = ax.text(0, cy, f"{cantidad:g}", ha="left", va="center",
                    fontsize=fontsize, fontweight="900", color=BLACK, zorder=50)
        w_num = ancho_texto(ax, t)
        total = w_num + hueco + 2 * radio * ASPECT
        if total <= ancho_max or fontsize <= 10:
            break
        t.remove()
        fontsize -= 1

    x0 = cx - total / 2
    t.set_x(x0)
    moneda(ax, x0 + w_num + hueco + radio * ASPECT, cy, valor, color, radio=radio, zbase=51)


def texto_con_sombra(ax, x, y, texto, **kwargs):
    # Mantiene números nítidos; no se aplica una sombra gruesa que reduzca la legibilidad.
    kwargs.setdefault("zorder", 30)
    return ax.text(x, y, texto, **kwargs)


def textura_papel(ax, seed=1234, n=1800):
    rng=np.random.default_rng(seed)
    ax.scatter(rng.random(n), rng.random(n),
               s=rng.uniform(0.1,1.0,n),
               c="#5B574F", alpha=rng.uniform(0.01,0.035,n),
               linewidths=0, zorder=1)
    for yy in rng.uniform(0,1,70):
        ax.plot([0,1],[yy,yy+rng.normal(0,0.0006)],
                color="#807B70", alpha=0.012, linewidth=0.45, zorder=1)

def textura_panel(ax,x,y,w,h,seed=42,n=80):
    rng=np.random.default_rng(seed)
    ax.scatter(x+rng.random(n)*w,y+rng.random(n)*h,
               s=rng.uniform(0.1,0.45,n),c="#FFFFFF",
               alpha=0.018,linewidths=0,zorder=3)

def sombra_panel(ax,x,y,w,h,alpha=0.055):
    ax.add_patch(Rectangle((x+0.002,y-0.003),w,h,
                           facecolor="#000000",edgecolor="none",
                           alpha=alpha,zorder=1))


def dibujar_celda(ax, x, y, w, h, casilla, valores, colores, indices):
    """Dibuja una celda limpia: cantidad + UNA ficha por denominación (datos del Excel)."""
    sombra_panel(ax, x, y, w, h, 0.045)
    ax.add_patch(Rectangle(
        (x, y), w, h,
        facecolor=CELL, edgecolor="#C7CFCA",
        linewidth=0.8, zorder=2
    ))
    textura_panel(ax, x, y, w, h, n=22)

    if casilla["estado"] == "NO ES POSIBLE":
        ax.text(
            x + w/2, y + h*0.60, "NO ES POSIBLE",
            ha="center", va="center", fontsize=12.5,
            fontweight="900", color="#A42A24", zorder=30
        )
        ax.text(
            x + w/2, y + h*0.27, "—",
            ha="center", va="center", fontsize=14,
            fontweight="bold", color=BLACK, zorder=30
        )
        return

    # Cantidades tal cual vienen del Excel, solo de los tipos de esta columna.
    partes = [
        (casilla["fichas"][idx], valores[idx], colores[idx])
        for idx in indices
        if casilla["fichas"][idx] > 0
    ]

    main_y = y + h*0.63

    if len(partes) == 1:
        cantidad, valor, color = partes[0]
        # Número grande a la izquierda + ficha pequeña, centrados en la celda.
        cantidad_y_ficha(ax, x + w/2, main_y, cantidad, valor, color,
                         fontsize=23, radio=0.024, hueco=0.007, ancho_max=w*0.85)

    elif len(partes) == 2:
        # Dos grupos claramente separados, cada uno en su mitad de la celda.
        # El número queda SIEMPRE a la izquierda de su ficha, a la misma altura.
        centers = [x + w*0.25, x + w*0.75]
        for (cantidad, valor, color), center in zip(partes, centers):
            cantidad_y_ficha(ax, center, main_y, cantidad, valor, color,
                             fontsize=19, radio=0.0175, hueco=0.004, ancho_max=w*0.42)

        # Separador entre los dos grupos.
        ax.text(
            x + w*0.50, main_y, "+",
            ha="center", va="center",
            fontsize=13, fontweight="900",
            color="#596563", zorder=50
        )

    # Banda de devolución, siempre legible y separada de la combinación.
    ax.add_patch(Rectangle(
        (x + 0.025*w, y + 0.06*h),
        0.95*w, 0.25*h,
        facecolor="#DCE1DE",
        edgecolor="none", zorder=4
    ))
    ax.text(
        x + w/2, y + h*0.185,
        f"Devuelve {fmt_euros(casilla['devolucion'])}€",
        ha="center", va="center",
        fontsize=11.5, fontweight="bold",
        color="#1A1A1A", zorder=30
    )


def generar_cartel(excel_path, output_path, forzar=False):
    valores, bebidas, colores, importes, tabla, textos = leer_excel(excel_path)

    # Python no corrige nada: si el Excel marca casillas a revisar, avisa y para.
    revisar = casillas_a_revisar(tabla, importes)
    if revisar:
        lista = "\n".join(f"  - {imp:g}€ · {comb}" for imp, comb in revisar)
        if not forzar:
            raise ErrorExcel(
                f"El Excel tiene {len(revisar)} casilla(s) en REVISAR (hoja CAMBIOS):\n{lista}\n"
                "Corrígelas y guarda el Excel, o usa --forzar para generar el cartel igualmente."
            )
        print(f"AVISO: se genera con {len(revisar)} casilla(s) en REVISAR:\n{lista}")

    # Aspecto del cartel
    fig = plt.figure(figsize=(FIG_W, FIG_H), facecolor=CREAM)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.set_facecolor(CREAM)
    textura_papel(ax)

    dark = "#0B2D33"
    black = "#07161A"

    # Encabezado
    ax.text(0.025, 0.950, textos["titulo"],
            fontsize=29, fontweight="900", color=black, va="center")
    ax.text(0.025, 0.902, textos["subtitulo"],
            fontsize=10.5, color=black, va="center", fontweight="bold")

    # Leyenda superior limpia: ficha + descripción en una sola línea.
    # Todo queda por encima de la banda y sin solapamientos.
    ax.text(0.765, 0.955, "TIPOS DE FICHAS",
            fontsize=18, fontweight="900", color=black,
            ha="center", va="center")

    legend_x = [0.60, 0.765, 0.93]
    for i, xx in enumerate(legend_x):
        moneda(ax, xx, 0.913, valores[i], colores[i], radio=0.020)

        ax.text(
            xx, 0.880,
            f"{valores[i]:g}€  ·  {bebidas[i].upper()}",
            fontsize=9, fontweight="900", color=black,
            ha="center", va="center"
        )

    # Banda título tabla
    sombra_panel(ax,0.008,0.812,0.984,0.055,0.09)
    ax.add_patch(Rectangle((0.008, 0.812), 0.984, 0.055,
                           facecolor=dark, edgecolor="none", zorder=2))
    textura_panel(ax,0.008,0.812,0.984,0.055,n=260)
    ax.text(0.5, 0.839, textos["tabla"],
            ha="center", va="center", fontsize=19, fontweight="900",
            color="white")

    # Geometría tabla: las filas ocupan el espacio hasta el pie.
    left = 0.008
    top = 0.79
    label_w = 0.105
    grid_left = left + label_w
    grid_w = 0.984 - label_w
    col_w = grid_w / 6
    header_h = 0.095
    bottom = 0.083
    row_h = min(0.13, (top - header_h - bottom) / len(importes))

    headers = [
        f"SOLO {bebidas[0].upper()}\n(fichas de {valores[0]:g}€)",
        f"SOLO {bebidas[1].upper()}\n(fichas de {valores[1]:g}€)",
        f"SOLO {bebidas[2].upper()}\n(fichas de {valores[2]:g}€)",
        f"{bebidas[0].upper()} + {bebidas[1].upper()}\n({valores[0]:g}€ + {valores[1]:g}€)",
        f"{bebidas[0].upper()} + {bebidas[2].upper()}\n({valores[0]:g}€ + {valores[2]:g}€)",
        f"{bebidas[1].upper()} + {bebidas[2].upper()}\n({valores[1]:g}€ + {valores[2]:g}€)",
    ]

    # encabezado de columnas
    for c, htxt in enumerate(headers):
        x = grid_left + c*col_w
        sombra_panel(ax,x,top-header_h,col_w-0.003,header_h-0.004,0.08)
        ax.add_patch(Rectangle((x, top-header_h), col_w-0.003, header_h-0.004,
                               facecolor=dark, edgecolor="white", linewidth=1,zorder=2))
        textura_panel(ax,x,top-header_h,col_w-0.003,header_h-0.004,n=35)
        ax.text(x + col_w/2, top-header_h/2, htxt,
                ha="center", va="center", fontsize=9.2,
                fontweight="900", color="white", linespacing=1.25)

    # filas
    for r, importe in enumerate(importes):
        y_top = top - header_h - r*row_h

        sombra_panel(ax,left,y_top-row_h+0.002,label_w-0.004,row_h-0.004,0.08)
        ax.add_patch(Rectangle((left, y_top-row_h+0.002), label_w-0.004, row_h-0.004,
                               facecolor=dark, edgecolor="white", linewidth=1,zorder=2))
        textura_panel(ax,left,y_top-row_h+0.002,label_w-0.004,row_h-0.004,n=18)
        ax.text(left + label_w/2, y_top-row_h/2,
                f"CAMBIO DE\n{importe:g}€",
                ha="center", va="center", fontsize=13,
                fontweight="900", color="white")

        for c, casilla in enumerate(tabla[importe]):
            x = grid_left + c*col_w
            dibujar_celda(ax, x, y_top-row_h, col_w-0.003, row_h-0.004,
                          casilla, valores, colores, COMBINACIONES[c])

    # Pie
    sombra_panel(ax,0.008,0.018,0.984,0.045,0.09)
    ax.add_patch(Rectangle((0.008, 0.018), 0.984, 0.045,
                           facecolor=dark, edgecolor="none",zorder=2))
    textura_panel(ax,0.008,0.018,0.984,0.045,n=220)
    ax.text(0.5, 0.040, textos["pie"],
            ha="center", va="center", fontsize=12,
            fontweight="bold", color="white")

    # Guardado
    fig.savefig(output_path, dpi=260, facecolor=CREAM,
                bbox_inches="tight", pad_inches=0.05)
    plt.close(fig)


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    forzar = "--forzar" in sys.argv

    base = Path(__file__).resolve().parent
    excel = Path(args[0]) if len(args) >= 1 else base / "plantilla_cambio_fichas.xlsx"
    output = Path(args[1]) if len(args) >= 2 else base / "cartel_cambio_fichas.png"

    try:
        generar_cartel(excel, output, forzar=forzar)
    except ErrorExcel as e:
        print(f"ERROR: {e}")
        sys.exit(1)
    except PermissionError:
        print(f"ERROR: no se puede leer/escribir. ¿Está '{excel.name}' o el PNG abierto en otro programa?")
        sys.exit(1)
    print(f"Cartel generado: {output.resolve()}")
