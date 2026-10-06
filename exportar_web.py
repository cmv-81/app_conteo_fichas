#!/usr/bin/env python3
"""
Exporta los datos del Excel a la web (public/datos.js).

Uso:
    python exportar_web.py [plantilla_cambio_fichas.xlsx]

Lee la misma hoja DATOS_PYTHON que el generador del cartel, así que la web y el
cartel muestran siempre los mismos valores. Después de ejecutarlo, sube los
cambios de la carpeta public/ a GitHub (Vercel publica solo).
"""

import json
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

from generador_cartel_fichas import HOJA_DATOS, ErrorExcel, casillas_a_revisar, leer_excel

BASE = Path(__file__).resolve().parent
SALIDA = BASE / "public" / "datos.js"


def numero(v):
    """5.0 -> 5 ; 2.5 -> 2.5 ; '—' -> None"""
    if isinstance(v, (int, float)):
        f = float(v)
        return int(f) if f.is_integer() else f
    return None


def leer_precio_vaso(excel_path):
    """Clave 'vaso_precio' de DATOS_PYTHON. Si el Excel es antiguo y no la tiene, 1 €."""
    ws = openpyxl.load_workbook(excel_path, data_only=True)[HOJA_DATOS]
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value is None:
            break
        if str(ws.cell(r, 1).value).strip() == "vaso_precio":
            precio = numero(ws.cell(r, 2).value)
            if precio is None or precio < 0:
                raise ErrorExcel("El precio del vaso del Excel no es válido (CONFIGURACIÓN, celda B23).")
            return precio
    return 1


def exportar(excel_path):
    valores, bebidas, colores, importes, tabla, _ = leer_excel(excel_path)
    precio_vaso = leer_precio_vaso(excel_path)

    revisar = casillas_a_revisar(tabla, importes)
    if revisar:
        lista = "\n".join(f"  - {imp:g}€ · {comb}" for imp, comb in revisar)
        raise ErrorExcel(
            f"El Excel tiene {len(revisar)} casilla(s) en REVISAR (hoja CAMBIOS):\n{lista}\n"
            "Corrígelas y guarda el Excel antes de exportar a la web."
        )

    datos = {
        "generado": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "excel": Path(excel_path).name,
        "fichas": [
            {"valor": numero(v), "nombre": b, "color": c}
            for v, b, c in zip(valores, bebidas, colores)
        ],
        "vaso": {"precio": precio_vaso},
        "cartel": [
            {
                "importe": numero(importe),
                "casillas": [
                    {
                        "estado": cas["estado"],
                        "fichas": cas["fichas"],
                        "devolucion": numero(cas["devolucion"]),
                    }
                    for cas in tabla[importe]
                ],
            }
            for importe in importes
        ],
    }

    contenido = (
        "// Generado automáticamente por exportar_web.py a partir de "
        f"{datos['excel']}. NO EDITAR A MANO.\n"
        "window.FICHAS_DATOS = "
        + json.dumps(datos, ensure_ascii=False, indent=1)
        + ";\n"
    )
    SALIDA.parent.mkdir(exist_ok=True)
    SALIDA.write_text(contenido, encoding="utf-8")
    return datos


if __name__ == "__main__":
    excel = Path(sys.argv[1]) if len(sys.argv) > 1 else BASE / "plantilla_cambio_fichas.xlsx"
    try:
        datos = exportar(excel)
    except ErrorExcel as e:
        print(f"ERROR: {e}")
        sys.exit(1)
    except PermissionError:
        print(f"ERROR: no se puede leer '{excel.name}'. ¿Está abierto y bloqueado por otro programa?")
        sys.exit(1)
    fichas = ", ".join(f"{f['nombre']} {f['valor']}€" for f in datos["fichas"])
    print(f"Web actualizada: {SALIDA}")
    print(f"  Fichas: {fichas}")
    print(f"  Vaso: {datos['vaso']['precio']}€")
    print(f"  Importes del cartel: {', '.join(str(c['importe']) + '€' for c in datos['cartel'])}")
